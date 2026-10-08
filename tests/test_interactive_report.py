"""Small source-shaped fixtures; no scientific tools or calculations are invoked."""
import csv
import hashlib
import json
from pathlib import Path
import tempfile
import unittest

from openpyxl import Workbook

from dorado_workflow.reports.interactive.coordinates import normalize_interval
from dorado_workflow.reports.interactive.importers import import_results
from dorado_workflow.reports.interactive.models import Availability as A, Evidence, Fact, safe_id, validate
from dorado_workflow.reports.interactive.builder import build_report
from dorado_workflow.reports.interactive.dashboard import dashboard_data, motif_comparison


class ReportDataTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.addCleanup(self.tmp.cleanup)

    def csv(self, name, rows, headers=None):
        path = self.root / name
        path.parent.mkdir(parents=True, exist_ok=True)
        with path.open("w", newline="", encoding="utf-8") as handle:
            writer = csv.DictWriter(handle, fieldnames=headers or list(rows[0]))
            writer.writeheader()
            writer.writerows(rows)
        return path

    def raw(self, barcode="01", rows=None):
        return self.csv(f"nanotel/barcode{barcode}_summary.csv", rows if rows is not None else [self.row()])

    def row(self, rid="read_001"):
        return dict(sequence_ID=rid, sequence_length=500, Telomere_start=1,
                    Telomere_end=120, Telomere_length=120, telo_density=0.9)

    def metadata(self, **kwargs):
        path = self.root / "metadata.json"
        path.write_text(json.dumps(kwargs), encoding="utf-8")
        return path

    def test_parallel_worker_density_filenames_are_barcode_scoped(self):
        self.raw("01")
        self.raw("02")
        sample = dict(read_id="read_001", method="exact", start=1, end=100,
                      density=0.42, convention="one_based_inclusive", space="nanotel_analyzed_read")
        self.csv("nanotel/barcode01/report_details/read131_density.csv", [sample])
        run = import_results(self.root)
        self.assertEqual(run.barcodes[0].reads[0].tracks["density"].status, A.AVAILABLE)
        self.assertNotEqual(run.barcodes[1].reads[0].tracks["density"].status, A.AVAILABLE)

    def test_motif_comparison_counts_hits_with_explicit_scope(self):
        self.raw()
        run=import_results(self.root)
        barcode=run.barcodes[0];read=barcode.reads[0]
        barcode.metrics["filtered_length_records"]=Fact(A.AVAILABLE,[{"read_id":read.id,"length_bp":120}])
        read.tracks["motifs"]=Fact(A.AVAILABLE,{"methods":["tvr_inclusive"],"matches":[
            dict(method="tvr_inclusive",motif="TCAGGG",role="tvr",start=6,end=12),
            dict(method="tvr_inclusive",motif="TCAGGG",role="tvr",start=200,end=206)]})
        entry=motif_comparison(run)[0]
        self.assertTrue(entry["available"])
        self.assertEqual(entry["telomere"],{"TCAGGG":1})
        self.assertEqual(entry["whole_read"],{"TCAGGG":2})
        read.tracks["motifs"]=Fact(A.NOT_EXPORTED)
        self.assertFalse(motif_comparison(run)[0]["available"])

    def test_motif_matches_zero_invalid_and_barcode_scope(self):
        self.raw("01")
        self.raw("02")
        marker=dict(read_id="read_001",method="tvr_inclusive",record_type="search_complete",
            motif="",matched_sequence="",start="",end="",role="",max_mismatch="",
            convention="one_based_inclusive",space="nanotel_analyzed_read")
        hit={**marker,"record_type":"match","motif":"TCAGGG","matched_sequence":"TCAGGG",
             "start":7,"end":12,"role":"tvr","max_mismatch":0}
        path="nanotel/barcode01/report_details/read1_motifs.csv"
        self.csv(path,[marker,hit])
        run=import_results(self.root)
        track=run.barcodes[0].reads[0].tracks["motifs"]
        self.assertEqual(track.status,A.AVAILABLE)
        self.assertEqual(track.value["matches"][0]["start"],6)
        self.assertEqual(track.value["matches"][0]["motif"],"TCAGGG")
        self.assertNotEqual(run.barcodes[1].reads[0].tracks["motifs"].status,A.AVAILABLE)
        self.csv(path,[marker])
        self.assertEqual(import_results(self.root).barcodes[0].reads[0].tracks["motifs"].value["matches"],[])
        for rows in [[hit],[marker,{**hit,"end":999}],[marker,{**hit,"motif":"<script>"}]]:
            self.csv(path,rows)
            self.assertEqual(import_results(self.root).barcodes[0].reads[0].tracks["motifs"].status,A.INVALID)

    def test_exported_density_samples_and_invalid_values(self):
        self.raw()
        sample = dict(read_id="read_001", method="exact", start=1, end=100,
                      density=0.42, convention="one_based_inclusive", space="nanotel_analyzed_read")
        name = "nanotel/report_details/barcode01_read1_density.csv"
        self.csv(name, [sample])
        run = import_results(self.root)
        track = run.barcodes[0].reads[0].tracks["density"]
        self.assertEqual(track.status, A.AVAILABLE)
        self.assertEqual(track.value, [dict(method="exact", start=0, end=100, density=0.42)])
        self.assertTrue(track.evidence)
        for changes in [dict(density=1.1), dict(end=501), dict(convention="unknown"), dict(space="reference")]:
            self.csv(name, [{**sample, **changes}])
            self.assertEqual(import_results(self.root).barcodes[0].reads[0].tracks["density"].status, A.INVALID)
        self.csv(name, [sample, sample])
        self.assertEqual(import_results(self.root).barcodes[0].reads[0].tracks["density"].status, A.INVALID)

    def test_interval_conversion_and_no_orientation_guess(self):
        self.raw()
        run = import_results(self.root)
        call = run.barcodes[0].reads[0].calls["exact"]
        self.assertEqual(call["interval"].value, dict(start=0, end=120,
            convention="zero_based_half_open", space="nanotel_analyzed_read"))
        self.assertEqual(call["original_read_orientation"].status, A.NOT_EXPORTED)
        self.assertEqual(call["raw_interval"].value["start"], "1")

    def test_invalid_and_unknown_coordinates(self):
        for start, end in [(0, 5), (8, 2), (1.5, 8), (1, 12), ("inf", 10)]:
            interval = normalize_interval(start, end, convention="one_based_inclusive",
                space="reference", evidence=[], length=10)
            self.assertEqual(interval["normalized"].status, A.INVALID)
        result = normalize_interval(1, 6, convention="unknown", space="reference", evidence=[])
        self.assertEqual(result["normalized"].status, A.NOT_EXPORTED)
        result = normalize_interval(0, 6, convention="zero_based_half_open", space="reference", evidence=[])
        self.assertEqual(result["normalized"].value["start"], 0)

    def test_missing_zero_invalid_are_distinct(self):
        row = self.row(); row["telo_density"] = 0
        self.raw(rows=[row])
        read = import_results(self.root).barcodes[0].reads[0]
        self.assertEqual(read.calls["exact"]["density"].value, 0)
        self.assertEqual(read.calls["mismatch"]["density"].status, A.NOT_EXPORTED)
        row["telo_density"] = "Infinity"
        self.raw(rows=[row])
        self.assertEqual(import_results(self.root).barcodes[0].reads[0].calls["exact"]["density"].status, A.INVALID)

    def test_sources_hash_rows_and_input_unchanged(self):
        path = self.raw(); before = path.read_bytes()
        run = import_results(self.root)
        self.assertEqual(run.sources[0].sha256, hashlib.sha256(before).hexdigest())
        self.assertEqual(run.barcodes[0].reads[0].evidence[0].row, 2)
        self.assertEqual(path.read_bytes(), before)
        validate(run)

    def test_no_filtering_or_call_preference(self):
        row = self.row(); row["Telomere_length_mismatch"] = 300
        row["Telomere_length_mismatch_tvr"] = 400
        self.raw(rows=[row, self.row("absent")])
        self.csv("nanotel/filtered_summary_barcode01.csv", [{"read_id": "read_001"}])
        run = import_results(self.root)
        first, absent = run.barcodes[0].reads
        self.assertEqual(first.filtering.value, "present_in_filtered_output")
        self.assertIsNone(absent.filtering.value)
        self.assertEqual([first.calls[k]["reported_length"].value for k in ("exact", "mismatch", "tvr_inclusive")], [120, 300, 400])

    def test_duplicate_read_ids_and_ambiguous_mapping(self):
        self.raw(rows=[self.row("same a"), self.row("same b"), self.row("same a")])
        self.csv("mapping/mappedbarcode01_combined.csv", [{"read_id_clean": "same", "alignment_genome": "chr1Head"}])
        barcode = import_results(self.root).barcodes[0]
        self.assertEqual(len({r.id for r in barcode.reads}), 3)
        self.assertTrue(barcode.warnings)
        self.assertTrue(all(r.alignments.value is None for r in barcode.reads))
        self.assertTrue(all(r.alignments.status == A.INVALID for r in barcode.reads))

    def test_safe_stable_ids_across_barcodes_and_hostile_strings(self):
        hostile = '../../a/#?<script>alert("x")</script> שלום'
        self.raw("01", [self.row(hostile)]); self.raw("02", [self.row(hostile)])
        first = import_results(self.root); second = import_results(self.root)
        ids = [b.reads[0].id for b in first.barcodes]
        self.assertEqual(ids, [b.reads[0].id for b in second.barcodes])
        self.assertNotEqual(*ids)
        self.assertRegex(ids[0], r"^r-[a-f0-9]{64}$")
        self.assertNotEqual(safe_id("r", "a/b", "c"), safe_id("r", "a", "b/c"))

    def test_summary_only_and_explicit_stages(self):
        self.raw()
        meta = self.metadata(stages={"summary_only": True, "mapping": False, "methylation": True})
        read = import_results(self.root, meta).barcodes[0].reads[0]
        self.assertEqual(read.tracks["density"].status, A.NOT_REQUESTED)
        self.assertEqual(read.alignments.status, A.NOT_REQUESTED)
        self.assertEqual(read.tracks["methylation"].status, A.MISSING)

    def test_mapping_preserves_multiple_alignments_and_unknown_arm(self):
        self.raw()
        row = dict(read_id_clean="read_001", alignment_genome="chr1Head", alignment_genome_start=1,
                   alignment_genome_end=200, alignment_direction="-", alignment_mapq=60)
        self.csv("mapping/mappedbarcode01_combined.csv", [row, dict(row, alignment_mapq=255)])
        run = import_results(self.root)
        mappings = run.barcodes[0].reads[0].alignments.value
        self.assertEqual(len(mappings), 2)
        self.assertIsNone(mappings[0]["normalized"].value)
        self.assertIsNone(mappings[0]["arm"].value)
        self.assertEqual(mappings[1]["mapq"].status, A.NOT_APPLICABLE)
        meta = self.metadata(mapping_coordinate_convention="one_based_inclusive")
        normalized = import_results(self.root, meta).barcodes[0].reads[0].alignments.value[0]["normalized"].value
        self.assertEqual((normalized["start"], normalized["end"]), (0, 200))
        meta = self.metadata(mapping_coordinate_convention="zero_based_half_open")
        normalized = import_results(self.root, meta).barcodes[0].reads[0].alignments.value[0]["normalized"].value
        self.assertEqual((normalized["start"], normalized["end"]), (1, 200))

    def test_workbook_fallback_empty_sheet_and_stats(self):
        book = Workbook(); stats = book.active; stats.title = "NanoTel Statistics"
        stats.append(["barcode", "amount_of_telomeres", "median_telomere_length"])
        stats.append(["barcode01", 0, None])
        raw = book.create_sheet("barcode01"); raw.append(list(self.row())); raw.append(list(self.row().values()))
        filtered = book.create_sheet("filtered_barcode01"); filtered.append(["read_id"])
        book.save(self.root / "nanotel_summary.xlsx"); book.close()
        run = import_results(self.root)
        self.assertEqual(run.barcodes[0].metrics["amount_of_telomeres"].value, 0)
        self.assertIsNone(run.barcodes[0].reads[0].filtering.value)
        self.assertEqual(run.barcodes[0].reads[0].evidence[0].row, 2)
        self.assertTrue(any(s.sheet == "barcode01" for s in run.sources))

    def test_empty_missing_and_duplicate_sources(self):
        run = import_results(self.root)
        self.assertTrue(run.warnings); self.assertIsNone(run.metrics["barcode_count"].value)
        self.csv("barcode01_summary.csv", [], headers=list(self.row()))
        self.assertEqual(import_results(self.root).barcodes[0].metrics["imported_read_records"].value, 0)
        self.raw()
        with self.assertRaisesRegex(ValueError, "duplicate summary"):
            import_results(self.root)

    def test_validation_and_reported_length_disagreement(self):
        row = self.row(); row["Telomere_length"] = 119
        self.raw(rows=[row])
        run = import_results(self.root)
        self.assertTrue(run.barcodes[0].reads[0].warnings)
        self.assertEqual(run.barcodes[0].reads[0].calls["exact"]["reported_length"].value,119)
        run.metrics["bad"] = Fact(A.MISSING, 0, "missing")
        with self.assertRaises(ValueError): validate(run)
        del run.metrics["bad"]
        run.barcodes[0].reads[0].barcode_id = "bad"
        with self.assertRaises(ValueError): validate(run)

    def test_offline_build_and_safe_payload(self):
        self.raw(rows=[self.row('</script><img src=x onerror=alert(1)>')])
        run = import_results(self.root)
        path = build_report(run, self.root / "report")
        self.assertTrue(path.is_file())
        chunk = path.parent / "assets/data" / (run.barcodes[0].id + ".js")
        self.assertNotIn("</script>", chunk.read_text())
        self.assertIn("\\u003c", chunk.read_text())
        stored = json.loads((path.parent / "report-data.json").read_text())
        self.assertEqual(stored["barcodes"][0]["reads"][0]["original_id"], run.barcodes[0].reads[0].original_id)
        javascript = (path.parent / "assets/report.js").read_text()
        self.assertNotIn("fetch(", javascript)
        self.assertNotIn("innerHTML", javascript)
        self.assertNotIn("https://", path.read_text())
        with self.assertRaises(FileExistsError): build_report(run, path.parent)
        # Reimporting results does not consume the generated report's data.
        self.assertEqual(len(import_results(self.root).barcodes[0].reads), 1)

    def test_cli_and_validation_errors(self):
        from contextlib import redirect_stdout
        from io import StringIO
        from dorado_workflow.reports.interactive.__main__ import main
        self.raw()
        with redirect_stdout(StringIO()) as output:
            self.assertEqual(main([str(self.root)]), 0)
        self.assertTrue(output.getvalue().startswith("file:///"))
        self.assertTrue((self.root / "report/index.html").is_file())
        with self.assertRaisesRegex(ValueError, "booleans"):
            import_results(self.root, self.metadata(stages={"mapping": "false"}))
        run = import_results(self.root)
        run.barcodes[0].reads[0].calls["exact"]["interval"].value["end"] = -1
        with self.assertRaisesRegex(ValueError, "normalized interval"):
            validate(run)

    def test_csv_precedence_and_filtered_only_sources(self):
        self.csv("filtered_summary_barcode01.csv", [self.row()])
        run = import_results(self.root)
        self.assertEqual(run.barcodes[0].metrics["imported_read_records"].population, "filtered_summary_records")
        self.assertEqual(run.barcodes[0].reads[0].filtering.value, "present_in_filtered_output")
        self.raw()
        book = Workbook(); book.active.title = "barcode01"
        book.active.append(list(self.row())); book.active.append(list(self.row("workbook_only").values()))
        book.save(self.root / "nanotel_summary.xlsx"); book.close()
        run = import_results(self.root)
        self.assertEqual([r.original_id for r in run.barcodes[0].reads], ["read_001"])
        self.assertTrue(any("CSV is authoritative" in message for message in run.warnings))

    def test_dashboard_aggregates_filtered_lengths_not_barcode_medians(self):
        stats = []
        for barcode, lengths in (("01", [1000, 2000, 3000]), ("02", [10000])):
            rows = [dict(self.row(f"read{i}"), sequence_length=20000, Telomere_length_mismatch=value) for i, value in enumerate(lengths)]
            self.raw(barcode, rows)
            self.csv(f"nanotel/filtered_summaryBARCODE{barcode}.csv", rows)
            stats.append({"barcode": f"barcode{barcode}", "amount_of_telomeres": len(rows), "median_telomere_length": lengths[len(lengths)//2]})
            logs = self.root / "logs"; logs.mkdir(exist_ok=True)
            (logs/f"nanotel_barcode{barcode}.log").write_text(f"Total reads in sample: 100\nNumber of reads which identified as Telomeric: {len(rows)}\nWork ended at: now\n")
        self.csv("nanotel/nanotel_summary_statistics.csv", stats)
        report_dir = self.root / "reports"; report_dir.mkdir()
        (report_dir/"r_pipeline_config.json").write_text(json.dumps({"run_mapping_analysis":False,"nanotel_analysis":{"summary_only":True}}))
        run = import_results(self.root)
        dashboard = dashboard_data(run)
        self.assertEqual(dashboard["metrics"]["total_reads"]["value"], 200)
        self.assertEqual(dashboard["metrics"]["telomeric_reads"]["value"], 4)
        self.assertEqual(dashboard["metrics"]["telomeric_percentage"]["value"], 2)
        self.assertEqual(dashboard["metrics"]["median_telomere_length"]["value"], 2500)
        self.assertEqual(run.stages["mapping"].value, False)
        self.assertEqual(run.stages["summary_only"].value, True)
        self.assertTrue(any(s.format=="log" for s in run.sources))
        (self.root/"logs/nanotel_barcode02.log").unlink()
        self.assertIsNone(dashboard_data(import_results(self.root))["metrics"]["total_reads"]["value"])

    def test_dashboard_rejects_partial_filtered_distribution(self):
        self.raw()
        self.csv("filtered_summaryBARCODE01.csv", [dict(self.row(), Telomere_length_mismatch=120)])
        self.csv("nanotel_summary_statistics.csv", [{"barcode":"barcode01","amount_of_telomeres":2}])
        run=import_results(self.root)
        self.assertEqual(run.barcodes[0].metrics["filtered_length_records"].status,A.INVALID)
        self.assertIsNone(dashboard_data(run)["metrics"]["median_telomere_length"]["value"])

    def test_dashboard_reference_dedup_and_cleaned_read_join(self):
        self.raw(rows=[self.row("read_001 runid=original")])
        self.csv("filtered_summaryBARCODE01.csv", [{"read_id":"read_001","Telomere_length_mismatch":120}])
        self.csv("nanotel_summary_statistics.csv", [{"barcode":"barcode01","amount_of_telomeres":1}])
        mapping={"read_id_clean":"read_001","alignment_genome":"chr1_Tail","alignment_mapq":60}
        self.csv("mapping/mappedbarcode01_combined.csv",[mapping,mapping,dict(mapping,alignment_genome="chr2_Tail")])
        run=import_results(self.root)
        records=run.barcodes[0].metrics["filtered_length_records"].value
        self.assertEqual(records[0]["read_id"],run.barcodes[0].reads[0].id)
        cells=dashboard_data(run)["reference_cells"]
        self.assertEqual(len(cells),2)
        self.assertTrue(all(c["read_count"]==1 and c["median_bp"]==120 for c in cells))


class OfflineBrowserTests(unittest.TestCase):
    """Exercise the generated local files in Qt's Chromium engine, offline."""

    @classmethod
    def setUpClass(cls):
        import os
        os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
        try:
            from PySide6.QtWidgets import QApplication
            from PySide6.QtWebEngineWidgets import QWebEngineView
            from PySide6.QtWebEngineCore import QWebEngineUrlRequestInterceptor
        except ImportError:
            raise unittest.SkipTest("Qt WebEngine unavailable; install GUI dependencies for browser tests")
        cls.app = QApplication.instance() or QApplication([])
        cls.View = QWebEngineView

        class OfflineInterceptor(QWebEngineUrlRequestInterceptor):
            def __init__(self, parent):
                super().__init__(parent)
                self.requests = []

            def interceptRequest(self, info):
                scheme = info.requestUrl().scheme()
                self.requests.append(scheme)
                if scheme not in {"file", "data", "about"}:
                    info.block(True)
        cls.Interceptor = OfflineInterceptor

    def evaluate(self, expression):
        from PySide6.QtCore import QEventLoop, QTimer
        loop = QEventLoop(); result = []
        self.view.page().runJavaScript("JSON.stringify(" + expression + ")", lambda value: (result.append(value), loop.quit()))
        timer = QTimer(); timer.setSingleShot(True); timer.timeout.connect(loop.quit); timer.start(15000)
        loop.exec()
        self.assertTrue(result, "JavaScript callback timed out")
        return json.loads(result[0]) if result[0] else None

    def wait_text(self, text):
        import time
        from PySide6.QtTest import QTest
        deadline = time.monotonic() + 10
        while time.monotonic() < deadline:
            actual = self.evaluate("document.body ? document.body.innerText : ''") or ""
            if text in actual:
                return actual
            QTest.qWait(75)
        self.fail(f"Expected {text!r}; rendered: {actual[:1200]}")

    def load_page(self, action):
        from PySide6.QtCore import QEventLoop, QTimer
        loop = QEventLoop(); finished = []
        def done(ok):
            finished.append(ok); loop.quit()
        self.view.loadFinished.connect(done)
        timer = QTimer(); timer.setSingleShot(True); timer.timeout.connect(loop.quit); timer.start(20000)
        action(); loop.exec()
        self.view.loadFinished.disconnect(done)
        self.assertEqual(finished, [True], "Local page failed to load")

    def click(self, selector):
        from PySide6.QtCore import QPoint, Qt
        from PySide6.QtTest import QTest
        rect = self.evaluate("(() => { const e=document.querySelector(" + json.dumps(selector) + "); e.scrollIntoView({block:'center'}); const r=e.getClientRects()[0]; return {x:r.x+r.width/2,y:r.y+r.height/2}; })()")
        QTest.mouseClick(self.view.focusProxy(), Qt.LeftButton, pos=QPoint(round(rect["x"]), round(rect["y"])))

    def capture(self, name):
        """Optional local review screenshots; normal test runs leave no artifacts."""
        import os
        directory = os.environ.get("NANOTEL_REPORT_SCREENSHOTS")
        if directory:
            from PySide6.QtTest import QTest
            QTest.qWait(150)
            path = Path(directory); path.mkdir(parents=True, exist_ok=True)
            self.view.grab().save(str(path / (name + ".png")))

    def test_file_navigation_reload_history_missing_data_and_no_network(self):
        from PySide6.QtCore import QUrl
        import shutil
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            rid = 'read_001 </script><img src=x onerror=alert(1)> שלום'
            with (root / "barcode01_summary.csv").open("w", newline="", encoding="utf-8") as handle:
                writer = csv.writer(handle)
                writer.writerow(["sequence_ID", "sequence_length", "Telomere_start", "Telomere_end", "Telomere_length"])
                writer.writerow([rid, 500, 1, 120, 120])
            with (root / "filtered_summaryBARCODE01.csv").open("w", newline="", encoding="utf-8") as handle:
                writer=csv.writer(handle); writer.writerow(["read_id","Telomere_length_mismatch"]); writer.writerow([rid,120])
            (root/"nanotel_summary_statistics.csv").write_text("barcode,amount_of_telomeres,median_telomere_length\nbarcode01,1,120\n")
            run = import_results(root)
            index = build_report(run, root / "report")
            b = run.barcodes[0]; read = b.reads[0]
            self.view = self.View()
            self.view.resize(1360, 900)
            interceptor = self.Interceptor(self.view)
            self.view.page().profile().setUrlRequestInterceptor(interceptor)
            try:
                self.view.show()
                self.load_page(lambda: self.view.load(QUrl.fromLocalFile(str(index))))
                self.wait_text("Key Statistics")
                from PySide6.QtTest import QTest
                for _ in range(100):
                    if self.evaluate("document.querySelectorAll('.plot[data-ready=true]').length") >= 2:
                        break
                    QTest.qWait(75)
                self.assertGreaterEqual(self.evaluate("document.querySelectorAll('.plot[data-ready=true]').length"),2)
                self.capture("run")
                self.click(".barlayer .point path")
                self.wait_text("Barcode statistics")
                self.capture("barcode")
                self.assertIn("Telomere length distribution", self.evaluate("document.body.textContent"))
                self.click("#view table a")
                text = self.wait_text("Source-reported telomere calls")
                self.capture("read")
                self.assertIn(rid, self.evaluate("document.body.textContent"))
                self.assertEqual(self.evaluate("document.querySelector('h1').textContent"), "Read: read_001")
                for _ in range(100):
                    if self.evaluate("document.querySelectorAll('.plot[data-ready=true]').length"):
                        break
                    QTest.qWait(75)
                self.assertEqual(self.evaluate("document.querySelectorAll('.plot[data-ready=true]').length"), 1)
                self.assertEqual(self.evaluate("document.querySelector('.js-plotly-plot').data.at(-1).x"), [1, 120])
                text = self.evaluate("document.body.textContent")
                self.assertIn("1–120 bp", text)
                self.assertIn("Source provenance", text)
                self.assertEqual(self.evaluate("document.querySelectorAll('img').length"), 0)
                fragment = f"#/barcode/{b.id}/read/{read.id}"
                self.assertEqual(self.evaluate("location.hash"), fragment)
                self.load_page(self.view.reload); self.wait_text("Source-reported telomere calls")
                self.view.back(); self.wait_text("Barcode statistics")
                self.view.forward(); self.wait_text("Source-reported telomere calls")
                self.evaluate("(location.hash='#/barcode/b-unknown', true)")
                self.wait_text("Unknown report address")
                self.evaluate("(document.querySelector('#view a').click(), true)")
                self.wait_text("Key Statistics")
                wrong_read = f"#/barcode/{b.id}/read/r-" + "0" * 64
                self.evaluate("(location.hash=" + json.dumps(wrong_read) + ", true)")
                self.wait_text("This read is not present in this barcode")
                # New local folder tests portability, direct-link loading and missing chunks.
                moved = root / "moved report"
                shutil.copytree(index.parent, moved)
                url = QUrl.fromLocalFile(str(moved / "index.html")); url.setFragment(fragment[1:])
                self.load_page(lambda: self.view.load(url)); self.wait_text("Source-reported telomere calls")
                (moved / "assets/data" / (b.id + ".js")).unlink()
                self.load_page(self.view.reload); self.wait_text("Barcode data is missing")
                (moved / "assets/data/run.js").unlink()
                self.load_page(self.view.reload); self.wait_text("Run data is missing")
                self.assertTrue(interceptor.requests)
                self.assertFalse(set(interceptor.requests) - {"file", "data", "about"})
            finally:
                self.view.page().profile().setUrlRequestInterceptor(None)
                self.view.close(); self.view.deleteLater(); self.app.processEvents()


if __name__ == "__main__":
    unittest.main()
