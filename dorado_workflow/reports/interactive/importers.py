"""Adapters for existing NanoTel CSVs and nanotel_summary.xlsx.

    Never scan sequences, apply thresholds, select a preferred call, or classify
    absent records as failed. Ambiguous joins remain visible as warnings.
"""
import csv
import hashlib
import json
import math
from pathlib import Path
import re

from openpyxl import load_workbook

from ..nanotel_workbook import barcode_from_summary_filename
from .coordinates import normalize_interval
from .models import (Availability as A, Barcode, Evidence, Fact, Read, Run,
                     Source, safe_id, unavailable, validate)


class ResultImporter:
    def __init__(self, root, metadata_path=None):
        self.root = Path(root).resolve()
        if not self.root.is_dir():
            raise ValueError("Results directory does not exist")
        self.sources = {}
        self.warnings = []
        self.metadata = {}
        self.meta_evidence = []
        saved_config = self.root / "reports/r_pipeline_config.json"
        if saved_config.is_file():
            saved = json.loads(saved_config.read_text(encoding="utf-8-sig"))
            self.metadata = {"settings": saved, "stages": {}}
            for stage, key in (("mapping", "run_mapping_analysis"), ("methylation", "run_methylation_analysis")):
                if type(saved.get(key)) is bool:
                    self.metadata["stages"][stage] = saved[key]
            summary_only = saved.get("nanotel_analysis", {}).get("summary_only")
            if type(summary_only) is bool:
                self.metadata["stages"]["summary_only"] = summary_only
            self.meta_evidence.append(Evidence(self.source(saved_config, "json").id))
        if metadata_path:
            path = Path(metadata_path).resolve()
            supplied = json.loads(path.read_text(encoding="utf-8-sig"))
            if not isinstance(supplied, dict):
                raise ValueError("Metadata must be an object")
            self.metadata.update(supplied)
            if "settings" in self.metadata and not isinstance(self.metadata["settings"], dict):
                raise ValueError("Metadata settings must be an object")
            if "organism" in self.metadata and not isinstance(self.metadata["organism"], str):
                raise ValueError("Metadata organism must be text")
            stages = self.metadata.get("stages", {})
            if not isinstance(stages, dict) or any(type(v) is not bool for v in stages.values()):
                raise ValueError("Metadata stages must contain booleans")
            convention = self.metadata.get("mapping_coordinate_convention", "unknown")
            if convention not in {"unknown", "one_based_inclusive", "zero_based_half_open"}:
                raise ValueError("Unsupported mapping coordinate convention")
            self.meta_evidence.append(Evidence(self.source(path, "json").id))

    def source(self, path, fmt, sheet=None):
        path = Path(path)
        label = path.relative_to(self.root).as_posix() if path.is_relative_to(self.root) else path.name
        key = safe_id("s", label, sheet or "")
        if key not in self.sources:
            digest = hashlib.sha256(path.read_bytes()).hexdigest()
            self.sources[key] = Source(key, label, digest, fmt, sheet)
        return self.sources[key]

    def files(self, pattern):
        # A previously generated report is never an input to its next build.
        return sorted(p for p in self.root.rglob(pattern)
                      if "report" not in p.relative_to(self.root).parts and p.is_file())

    def csv_table(self, path):
        source = self.source(path, "csv")
        with path.open(encoding="utf-8-sig", newline="") as handle:
            reader = csv.DictReader(handle)
            headers = reader.fieldnames or []
            if len(headers) != len(set(headers)) or not headers:
                raise ValueError(f"Invalid/duplicate CSV columns: {path.name}")
            rows = []
            for number, row in enumerate(reader, 2):
                if None in row:
                    raise ValueError(f"Extra CSV columns: {path.name}:{number}")
                rows.append((row, Evidence(source.id, number)))
            return headers, rows, source

    def workbook_tables(self, path):
        workbook = load_workbook(path, read_only=True, data_only=True)
        tables = {}
        try:
            for sheet in workbook:
                if sheet.title != "NanoTel Statistics" and not re.fullmatch(r"(?:filtered_)?barcode\d+", sheet.title):
                    continue
                source = self.source(path, "xlsx", sheet.title)
                iterator = iter(sheet.iter_rows(values_only=True))
                headers = list(next(iterator, ()))
                if not headers or any(not isinstance(h, str) for h in headers) or len(headers) != len(set(headers)):
                    raise ValueError(f"Invalid workbook columns: {sheet.title}")
                rows = [(dict(zip(headers, values)), Evidence(source.id, number))
                        for number, values in enumerate(iterator, 2) if any(v is not None for v in values)]
                tables[sheet.title] = (headers, rows, source)
        finally:
            workbook.close()
        return tables

    def unavailable_stage(self, stage, reason):
        requested = self.metadata.get("stages", {}).get(stage)
        status = A.NOT_REQUESTED if requested is False else A.MISSING if requested is True else A.NOT_EXPORTED
        return unavailable(reason, status)

    def motif_details(self, barcode):
        """Read preserved R search hits, never search sequences or call regions."""
        paths = [p for p in self.files("*read*_motifs.csv")
                 if p.parent.name == "report_details" and
                 (p.name.startswith(f"{barcode.label}_read") or
                  (re.fullmatch(r"read\d+_motifs\.csv", p.name) and
                   p.parent.parent.name.lower() == barcode.label.lower()))]
        index = {}
        for read in barcode.reads:
            index.setdefault(read.original_id, []).append(read)
        records, completed, invalid, evidence = {}, {}, set(), {}
        for path in paths:
            _, rows, source = self.csv_table(path)
            recovery_path = path.parent / "motif_recovery.json"
            recovery = Evidence(self.source(recovery_path, "json").id) if recovery_path.is_file() else None
            for row, ev in rows:
                targets = index.get(row.get("read_id"), [])
                if len(targets) != 1:
                    barcode.warnings.append(f"{path.name}: unmatched or ambiguous motif read ID")
                    continue
                read = targets[0]
                method = row.get("method")
                evidence.setdefault(read.id, []).append(ev)
                if recovery and recovery not in evidence[read.id]:
                    evidence[read.id].append(recovery)
                if (method not in read.calls or row.get("convention") != "one_based_inclusive"
                        or row.get("space") != "nanotel_analyzed_read"):
                    invalid.add(read.id); continue
                if row.get("record_type") == "search_complete":
                    methods = completed.setdefault(read.id, set())
                    if method in methods:
                        invalid.add(read.id)
                    methods.add(method)
                    continue
                interval = normalize_interval(row.get("start"), row.get("end"),
                    convention=row.get("convention"), space=row.get("space"), evidence=[ev],
                    length=read.metrics["read_length"].value)["normalized"]
                motif, matched = row.get("motif", ""), row.get("matched_sequence", "")
                allowance = row.get("max_mismatch")
                if (row.get("record_type") != "match" or interval.status != A.AVAILABLE
                        or row.get("role") not in {"canonical", "tvr"}
                        or not re.fullmatch(r"[ACGTWSMKRYBDHVN]+", motif)
                        or not re.fullmatch(r"[ACGTWSMKRYBDHVN]+", matched)
                        or allowance not in {"0", "1"}
                        or len(matched) != interval.value["end"] - interval.value["start"]):
                    invalid.add(read.id); continue
                records.setdefault(read.id, []).append(dict(method=method, motif=motif,
                    matched_sequence=matched, role=row["role"], max_mismatch=int(allowance),
                    start=interval.value["start"], end=interval.value["end"]))
        for read in barcode.reads:
            if read.id not in evidence:
                continue
            hits = records.get(read.id, [])
            methods = completed.get(read.id, set())
            if read.id in invalid or not methods or any(h["method"] not in methods for h in hits):
                read.tracks["motifs"] = unavailable("Invalid or incomplete motif export", A.INVALID)
            else:
                read.tracks["motifs"] = Fact(A.AVAILABLE,
                    {"methods": sorted(methods), "matches": hits}, evidence=evidence[read.id])
                read.evidence.extend(Evidence(sid) for sid in {e.source_id for e in evidence[read.id]})

    def density_details(self, barcode):
        """Import exported plot samples; never reconstruct density from sequence."""
        # Parallel R workers may not inherit the global barcode filename prefix.
        # In that case the enclosing barcode directory supplies the identity.
        paths = [p for p in self.files("*read*_density.csv")
                 if p.parent.name == "report_details" and
                 (p.name.startswith(f"{barcode.label}_read") or
                  (re.fullmatch(r"read\d+_density\.csv", p.name) and
                   p.parent.parent.name.lower() == barcode.label.lower()))]
        index = {}
        for read in barcode.reads:
            index.setdefault(read.original_id, []).append(read)
        collected, invalid = {}, set()
        for path in paths:
            _, rows, source = self.csv_table(path)
            for row, ev in rows:
                targets = index.get(row.get("read_id"), [])
                if len(targets) != 1:
                    barcode.warnings.append(f"{path.name}: density read ID unmatched or ambiguous")
                    continue
                read = targets[0]
                interval = normalize_interval(row.get("start"), row.get("end"),
                    convention=row.get("convention"), space=row.get("space"), evidence=[ev])["normalized"]
                density = number(row, "density", ev, maximum=1)
                method = row.get("method")
                if (interval.status != A.AVAILABLE or density.status != A.AVAILABLE
                        or row.get("space") != "nanotel_analyzed_read"
                        or method not in read.calls
                        or (read.metrics["read_length"].status == A.AVAILABLE
                            and interval.value["end"] > read.metrics["read_length"].value)):
                    invalid.add(read.id)
                    continue
                collected.setdefault(read.id, []).append(({
                    "method": method, "start": interval.value["start"],
                    "end": interval.value["end"], "density": density.value}, ev))
        for read in barcode.reads:
            samples = collected.get(read.id, [])
            seen = set()
            for sample, _ in samples:
                key = (sample["method"], sample["start"])
                if key in seen:
                    invalid.add(read.id)
                seen.add(key)
            if read.id in invalid:
                read.tracks["density"] = unavailable("Invalid or duplicate exported density samples", A.INVALID)
            elif samples:
                read.tracks["density"] = Fact(A.AVAILABLE,
                    sorted([s for s, _ in samples], key=lambda s: (s["method"], s["start"])),
                    unit="fraction", evidence=[e for _, e in samples])
                read.evidence.extend(Evidence(sid) for sid in {e.source_id for _, e in samples})

    def run(self):
        raw, filtered, stats = {}, {}, None
        for pattern, target in (("barcode*_summary.csv", raw), ("filtered_summary*.csv", filtered)):
            for path in self.files(pattern):
                barcode = barcode_from_summary_filename(path.name)
                if barcode is None:
                    self.warnings.append(f"Unrecognized barcode filename: {path.name}")
                    continue
                if barcode in target:
                    raise ValueError(f"Ambiguous duplicate summary for {barcode}")
                target[barcode] = self.csv_table(path)
        stat_paths = self.files("nanotel_summary_statistics.csv")
        if len(stat_paths) > 1:
            raise ValueError("Ambiguous run statistics CSVs")
        if stat_paths:
            stats = self.csv_table(stat_paths[0])
        workbooks = self.files("nanotel_summary.xlsx")
        if len(workbooks) > 1:
            raise ValueError("Ambiguous NanoTel workbooks; select one results directory")
        if workbooks:
            tables = self.workbook_tables(workbooks[0])
            stats = stats or tables.get("NanoTel Statistics")
            for name, table in tables.items():
                if name.startswith("filtered_"):
                    filtered.setdefault(name[len("filtered_"):], table)
                elif name.startswith("barcode"):
                    if name in raw:
                        self.warnings.append(f"{name}: CSV is authoritative; workbook duplicate was not imported")
                    raw.setdefault(name, table)

        stat_rows = {}
        if stats:
            for row, ev in stats[1]:
                label = str(row.get("barcode") or "")
                barcode = barcode_from_summary_filename(label)
                if not barcode:
                    self.warnings.append(f"Unrecognized statistics barcode: {label}")
                    continue
                if barcode in stat_rows:
                    raise ValueError(f"Duplicate statistics for {barcode}")
                stat_rows[barcode] = (row, ev)

        self.mapping_paths = self.files("mapped*.csv")
        barcodes = []
        for label in sorted(set(raw) | set(filtered) | set(stat_rows)):
            b = Barcode(safe_id("b", label), label)
            table = raw.get(label) or filtered.get(label)
            if table:
                population = "detected_summary_records" if label in raw else "filtered_summary_records"
                b.evidence.append(Evidence(table[2].id))
                if not any(key in table[0] for key in ("sequence_ID", "read_id")):
                    raise ValueError(f"Missing read ID column for {label}")
                counts = {}
                for row, ev in table[1]:
                    read_id = identity(row)
                    if not read_id:
                        b.warnings.append(f"Skipped row {ev.row}: missing read identifier")
                        continue
                    ordinal = counts.get(read_id, 0)
                    counts[read_id] = ordinal + 1
                    read = self.read(row, ev, b.id, ordinal)
                    if population == "filtered_summary_records":
                        read.filtering = Fact(A.AVAILABLE, "present_in_filtered_output", evidence=[ev])
                    b.reads.append(read)
                duplicates = {name for name, count in counts.items() if count > 1}
                for read in b.reads:
                    if read.original_id in duplicates:
                        read.warnings.append("Duplicate source read ID; records retained separately; joins are ambiguous")
                b.metrics["imported_read_records"] = Fact(A.AVAILABLE, len(b.reads), unit="records",
                    population=population, origin="presentation_aggregate", evidence=b.evidence.copy())
                if duplicates:
                    b.warnings.append("Duplicate read IDs exist; record count is not a unique-molecule count")
            else:
                b.metrics["imported_read_records"] = unavailable("Statistics exist but read detail is missing", A.MISSING)
            if label in raw and label in filtered:
                members = {identity(row) for row, _ in filtered[label][1]}
                for read in b.reads:
                    if read.original_id in members and not read.warnings:
                        read.filtering = Fact(A.AVAILABLE, "present_in_filtered_output",
                            evidence=[Evidence(filtered[label][2].id)])
                # Absence is not proof of failure; workbook empty sheets may be synthesized.
            if label in stat_rows:
                row, ev = stat_rows[label]
                b.evidence.append(ev)
                for key, unit in (("amount_of_telomeres", "reads"), ("median_telomere_length", "bp"),
                                  ("mean_density", "fraction")):
                    b.metrics[key] = number(row, key, ev, unit=unit, population="post_filtered_telomeres",
                                            maximum=1 if unit == "fraction" else None)
            for key in ("total_reads", "telomeric_percentage", "tvr_percentage", "reads_with_tvrs"):
                b.metrics[key] = unavailable("Not explicitly exported in supported summary tables")
            for key in ("amount_of_telomeres", "median_telomere_length", "mean_density"):
                b.metrics.setdefault(key, unavailable("Post-filtered statistics table is unavailable"))
            self.filtered_lengths(b, filtered.get(label))
            self.log_counts(b)
            self.mapping(b)
            self.density_details(b)
            self.motif_details(b)
            barcodes.append(b)

        if not barcodes:
            self.warnings.append("No supported barcode summaries or statistics found; this is not proof of zero reads")
        stages = {key: Fact(A.AVAILABLE, value, evidence=self.meta_evidence)
                  for key, value in self.metadata.get("stages", {}).items()}
        for key in ("mapping", "methylation", "summary_only"):
            stages.setdefault(key, unavailable("Run stage was not explicitly recorded"))
        settings = Fact(A.AVAILABLE, self.metadata["settings"], evidence=self.meta_evidence) if self.metadata.get("settings") is not None else unavailable("Settings snapshot not supplied")
        organism = Fact(A.AVAILABLE, self.metadata["organism"], evidence=self.meta_evidence) if self.metadata.get("organism") else unavailable("Organism not recorded in summary files")
        if organism.status != A.AVAILABLE:
            found = []
            for path in self.files("*.log"):
                if path.name.startswith("nanotel_"):
                    continue
                for line_number, line in enumerate(path.read_text(encoding="utf-8", errors="replace").splitlines(), 1):
                    match = re.search(r" - INFO - Organism: ([A-Za-z]+)\s*$", line)
                    if match:
                        found.append((match.group(1), Evidence(self.source(path, "log").id, line_number)))
            if found and len({value for value, _ in found}) == 1:
                organism = Fact(A.AVAILABLE, found[0][0], evidence=[ev for _, ev in found])
        label = self.root.parent.name if self.root.name.lower() == "results" else self.root.name
        run = Run(safe_id("run", label), label, organism, settings, stages,
                  barcodes, list(self.sources.values()), {}, self.warnings)
        run.metrics["barcode_count"] = Fact(A.AVAILABLE, len(barcodes), unit="barcodes",
            population="imported_barcodes", origin="presentation_aggregate",
            evidence=[Evidence(s.id) for s in run.sources]) if run.sources else unavailable("No supported sources found", A.MISSING)
        for key in ("total_reads", "telomeric_reads", "telomeric_percentage", "median_telomere_length"):
            run.metrics[key] = unavailable("No source-reported run-wide statistic in supported exports; see barcode statistics")
        validate(run)
        return run

    def filtered_lengths(self, barcode, table):
        """Use already-filtered lengths; reject inconsistent partial distributions."""
        key = "filtered_length_records"
        barcode.metrics[key] = unavailable("Filtered per-read lengths were not exported")
        if not table:
            return
        rows = []; seen = set(); counts = {}
        cleaned = {}
        by_report_id = {read.id: read for read in barcode.reads}
        for read in barcode.reads:
            counts.setdefault(read.original_id, []).append(read.id)
            cleaned.setdefault(read.original_id.split()[0], []).append(read.id)
        for row, ev in table[1]:
            rid = identity(row)
            length = number(row, "Telomere_length_mismatch", ev, unit="bp")
            if not rid or rid in seen or length.status != A.AVAILABLE:
                barcode.metrics[key] = unavailable("Filtered lengths contain missing/duplicate IDs or invalid lengths", A.INVALID)
                return
            seen.add(rid)
            # R may strip FASTQ header annotations in filtered tables. Join the
            # first read-ID token only when it identifies exactly one molecule.
            matching = counts.get(rid, []) or cleaned.get(rid.split()[0], [])
            rows.append({"original_id": rid, "read_id": matching[0] if len(matching) == 1 else None, "length_bp": length.value})
            if len(matching) == 1:
                target = by_report_id[matching[0]]
                target.filtering = Fact(A.AVAILABLE, "present_in_filtered_output", evidence=[ev])
        reported = barcode.metrics.get("amount_of_telomeres")
        if reported and reported.status == A.AVAILABLE and reported.value != len(rows):
            barcode.metrics[key] = unavailable("Filtered row count disagrees with source summary", A.INVALID)
            return
        if not rows and (reported is None or reported.status != A.AVAILABLE):
            return  # Empty fallback sheets may be synthesized by the Excel writer.
        barcode.metrics[key] = Fact(A.AVAILABLE, rows, unit="bp", population="post_filtered_telomeres",
            evidence=[Evidence(table[2].id, fields=["read_id", "Telomere_length_mismatch"])])

    def log_counts(self, barcode):
        """Read exact labels emitted by NanoTel, not guessed read counts."""
        paths = [p for p in self.files("nanotel_barcode*.log") if barcode_from_summary_filename(p.name) == barcode.label]
        if len(paths) != 1:
            return
        path = paths[0]
        lines = path.read_text(encoding="utf-8", errors="replace").splitlines()
        if not any(line.startswith("Work ended at:") for line in lines):
            return
        for key, pattern in (("total_reads", r"Total reads in sample: (\d+)"),
                             ("detected_telomeric_reads", r"Number of reads which identified as Telomeric: (\d+)")):
            found = [(i, re.fullmatch(pattern, line.strip())) for i, line in enumerate(lines, 1)]
            found = [(i, match) for i, match in found if match]
            if len(found) == 1:
                row, match = found[0]
                barcode.metrics[key] = Fact(A.AVAILABLE, int(match.group(1)), unit="reads",
                    population="input_reads" if key == "total_reads" else "detected_telomeric_reads",
                    evidence=[Evidence(self.source(path, "log").id, row)])

    def read(self, row, ev, barcode_id, ordinal):
        rid = identity(row)
        metrics = {"read_length": number(row, "sequence_length", ev, unit="bp")}
        calls = {}
        warnings = []
        for method, suffix in (("exact", ""), ("mismatch", "_mismatch"), ("tvr_inclusive", "_mismatch_tvr")):
            keys = ["Telomere_start" + suffix, "Telomere_end" + suffix]
            evidence = [Evidence(ev.source_id, ev.row, keys)]
            length = metrics["read_length"].value
            interval = normalize_interval(row.get(keys[0]), row.get(keys[1]),
                convention="one_based_inclusive", space="nanotel_analyzed_read", evidence=evidence, length=length)
            reported = number(row, "Telomere_length" + suffix, ev, unit="bp")
            calls[method] = {"reported_length": reported,
                "density": number(row, "telo_density" + suffix, ev, unit="fraction", maximum=1),
                "raw_interval": interval["raw"], "interval": interval["normalized"],
                "original_read_orientation": unavailable("Source does not record reverse-complement/crop transformation")}
            if interval["normalized"].status == A.AVAILABLE and reported.status == A.AVAILABLE:
                coords = interval["normalized"].value
                if coords["end"] - coords["start"] != reported.value:
                    warnings.append(f"{method}: reported length differs from interval span; reported value retained")
        tracks = {key: unavailable("Detailed track is not imported in milestone 1")
                  for key in ("density", "motifs", "tvr_regions")}
        if self.metadata.get("stages", {}).get("summary_only") is True:
            tracks = {key: unavailable("Summary-only run: detailed tracks were not requested", A.NOT_REQUESTED) for key in tracks}
        tracks["methylation"] = self.unavailable_stage("methylation", "No molecule-level methylation export supplied")
        return Read(safe_id("r", barcode_id, rid, str(ordinal)), rid, barcode_id, metrics, calls,
            unavailable("Filtering membership is not established; absence does not mean failure"),
            self.unavailable_stage("mapping", "No unambiguous mapping record available"), tracks, [ev], warnings)

    def mapping(self, barcode):
        candidates = [p for p in self.mapping_paths
                      if barcode_from_summary_filename(p.name) == barcode.label]
        combined = [p for p in candidates if p.name.endswith("_combined.csv")]
        paths = combined or candidates
        if len(combined) > 1:
            barcode.warnings.append("Multiple combined mapping files; mapping join skipped")
            for read in barcode.reads:
                read.alignments = unavailable("Ambiguous combined mapping sources", A.INVALID)
            return
        if paths and self.metadata.get("stages", {}).get("mapping") is False:
            barcode.warnings.append("Mapping files exist although supplied metadata says mapping was disabled; source records are retained")
        index = {}
        for read in barcode.reads:
            index.setdefault(read.original_id.split()[0], []).append(read)
        matches = {}
        for path in paths:
            _, rows, source = self.csv_table(path)
            barcode.evidence.append(Evidence(source.id))
            for row, ev in rows:
                rid = str(row.get("read_id_clean") or row.get("read_id.x") or row.get("read_id") or "").split()
                targets = index.get(rid[0], []) if rid else []
                if len(targets) != 1:
                    barcode.warnings.append(f"Mapping row {ev.row}: unmatched or ambiguous read ID; join skipped")
                    for target in targets:
                        target.alignments = unavailable("Cleaned read ID matches multiple source records; mapping join is ambiguous", A.INVALID)
                    continue
                interval = normalize_interval(row.get("alignment_genome_start"), row.get("alignment_genome_end"),
                    convention=self.metadata.get("mapping_coordinate_convention", "unknown"),
                    space="reference", evidence=[ev])
                alignment = {"reference": text_fact(row, "alignment_genome", ev),
                    "strand": text_fact(row, "alignment_direction", ev),
                    "mapq": number(row, "alignment_mapq", ev, maximum=255),
                    "chromosome": unavailable("Reference-to-chromosome annotation not supplied"),
                    "arm": unavailable("Head/Tail does not establish a chromosome arm"),
                    **interval}
                if alignment["mapq"].value == 255:
                    alignment["mapq"] = unavailable("SAM MAPQ 255 means mapping quality unavailable", A.NOT_APPLICABLE)
                matches.setdefault(targets[0].id, []).append((alignment, ev))
        for read in barcode.reads:
            if read.id in matches:
                pairs = matches[read.id]
                read.alignments = Fact(A.AVAILABLE, [a for a, _ in pairs], evidence=[e for _, e in pairs])


def identity(row):
    value = row.get("sequence_ID") or row.get("read_id")
    return str(value).strip() if value is not None else ""


def number(row, key, ev, *, unit=None, population=None, maximum=None):
    value = row.get(key)
    if value is None or str(value).strip() in {"", "NA", "NaN", "nan"}:
        return unavailable(f"{key} not exported or unavailable")
    try:
        parsed = float(value)
        if isinstance(value, bool) or not math.isfinite(parsed) or parsed < 0 or (maximum is not None and parsed > maximum):
            raise ValueError()
        if unit in {"bp", "reads"} and not parsed.is_integer():
            # Summary medians can be fractional; individual lengths cannot.
            if population is None or unit == "reads":
                raise ValueError()
        return Fact(A.AVAILABLE, int(parsed) if parsed.is_integer() else parsed, unit=unit,
                    population=population, evidence=[Evidence(ev.source_id, ev.row, [key])])
    except (ValueError, TypeError, OverflowError):
        return unavailable(f"Invalid source value for {key}: {value!r}", A.INVALID)


def text_fact(row, key, ev):
    value = row.get(key)
    return Fact(A.AVAILABLE, str(value), evidence=[Evidence(ev.source_id, ev.row, [key])]) if value not in (None, "") else unavailable(f"{key} missing")


def import_results(root, metadata_path=None) -> Run:
    """Read one results directory without modifying any scientific artifact."""
    return ResultImporter(root, metadata_path).run()
