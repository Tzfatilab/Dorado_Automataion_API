"""Write NanoTel CSV results to Excel without changing analysis values."""
import csv
import os
import re
from pathlib import Path
from typing import Dict, List, Optional
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

def barcode_from_summary_filename(filename: str) -> Optional[str]:
    """Return a normalized barcode from a CSV filename, or None if absent."""
    match = re.search(r"barcode\s*0*(\d+)", filename, re.IGNORECASE)
    if not match:
        match = re.search(r"(?:^|[_-])bc\s*0*(\d+)", filename, re.IGNORECASE)
    return f"barcode{int(match.group(1)):02d}" if match else None


class NanoTelWorkbookWriter:
    """Package raw, filtered and run-level CSVs into one workbook.

    Source CSVs are retained; the workflow decides when cleanup is safe.
    """

    def __init__(self, source_dir: Path, results_dir: Path,
                 short_telomere_threshold_bp: int = 2000):
        """Use CSVs in source_dir and save the workbook in results_dir.

        The threshold labels an empty summary; it does not filter any reads.
        """
        self.source_dir = source_dir
        self.results_dir = results_dir
        self.short_telomere_threshold_bp = short_telomere_threshold_bp

    def write(self) -> Path:
        """Return the saved workbook path, preserving CSV values and sheet order.

        Raise if raw summaries are missing or duplicate barcode names exist.
        Replace the final workbook only after the temporary file is saved.
        """
        summary_csv = self.source_dir / "nanotel_summary_statistics.csv"

        workbook = Workbook()
        summary_sheet = workbook.active
        summary_sheet.title = "NanoTel Statistics"

        raw_files, filtered_files = self._find_barcode_detail_files()
        if not raw_files:
            raise FileNotFoundError(
                f"No barcode NanoTel summaries were found in {self.source_dir}"
            )

        if summary_csv.exists():
            self._populate_csv_sheet(summary_sheet, summary_csv)
        else:
            self._populate_empty_summary_sheet(
                summary_sheet,
                self.short_telomere_threshold_bp,
            )
        self._ensure_summary_barcodes(summary_sheet, list(raw_files))

        for barcode, detail_file in raw_files.items():
            sheet = workbook.create_sheet(self._safe_sheet_name(barcode, workbook.sheetnames))
            self._populate_csv_sheet(sheet, detail_file)

        for barcode in raw_files:
            preferred_name = f"filtered_{barcode}"
            sheet = workbook.create_sheet(
                self._safe_sheet_name(preferred_name, workbook.sheetnames)
            )
            detail_file = filtered_files.get(barcode)
            if detail_file is None:
                self._populate_empty_filtered_sheet(sheet, raw_files[barcode])
            else:
                self._populate_csv_sheet(sheet, detail_file)

        workbook_path = self.results_dir / "nanotel_summary.xlsx"
        temporary_path = workbook_path.with_suffix(".tmp.xlsx")
        workbook.save(temporary_path)
        os.replace(temporary_path, workbook_path)
        return workbook_path


    def _find_barcode_detail_files(self) -> tuple[Dict[str, Path], Dict[str, Path]]:
        """Return ordered raw and filtered CSV sources, rejecting duplicates."""
        return (
            self._index_detail_files("barcode*_summary.csv", "NanoTel"),
            self._index_detail_files("filtered_summary*.csv", "filtered"),
        )

    def _index_detail_files(self, pattern: str, description: str) -> Dict[str, Path]:
        """Index matching CSVs by barcode without silently replacing a sample."""
        details: Dict[str, Path] = {}
        for path in sorted(self.source_dir.glob(pattern)):
            barcode = barcode_from_summary_filename(path.name)
            if not barcode:
                continue
            if barcode in details:
                raise ValueError(
                    f"Multiple {description} summaries resolve to {barcode}: "
                    f"{details[barcode].name}, {path.name}"
                )
            details[barcode] = path
        return dict(sorted(details.items()))

    @staticmethod
    def _safe_sheet_name(barcode: str, existing_names: List[str]) -> str:
        base = re.sub(r"[\\/*?:\[\]]", "_", barcode)[:31] or "Barcode"
        name = base
        suffix = 2
        while name in existing_names:
            tail = f"_{suffix}"
            name = f"{base[:31 - len(tail)]}{tail}"
            suffix += 1
        return name


    @staticmethod
    def _parse_csv_value(value: str):
        """Preserve identifiers as text while storing numeric results as numbers."""
        text = value.strip()
        if text == "":
            return None
        if re.fullmatch(r"[-+]?\d+", text):
            return int(text)
        if re.fullmatch(r"[-+]?(?:\d+\.\d*|\.\d+)(?:[eE][-+]?\d+)?", text):
            return float(text)
        return text


    def _populate_csv_sheet(self, sheet, csv_path: Path) -> None:
        """Copy a CSV into a readable, filterable worksheet."""
        with csv_path.open("r", encoding="utf-8-sig", newline="") as handle:
            reader = csv.reader(handle)
            column_widths: List[int] = []
            text_columns = set()
            row_count = 0
            for row_index, row in enumerate(reader, start=1):
                row_count = row_index
                while len(column_widths) < len(row):
                    column_widths.append(0)
                if row_index == 1:
                    text_columns = {
                        index
                        for index, name in enumerate(row)
                        if name.strip().lower() in {
                            "barcode", "read_id", "sequence_id", "read_name"
                        }
                        or name.strip().lower().endswith("_id")
                    }
                for column_index, value in enumerate(row, start=1):
                    parsed_value = (
                        value
                        if row_index == 1 or column_index - 1 in text_columns
                        else self._parse_csv_value(value)
                    )
                    cell = sheet.cell(
                        row=row_index,
                        column=column_index,
                        value=parsed_value,
                    )
                    # Read IDs and other CSV text beginning with '=' are literal
                    # data; openpyxl otherwise interprets them as Excel formulas.
                    if isinstance(parsed_value, str) and parsed_value.startswith("="):
                        cell.data_type = "s"
                        cell.number_format = "@"
                    if row_index > 1 and isinstance(cell.value, (int, float)):
                        cell.number_format = "#,##0.###"
                    if row_index <= 250:
                        column_widths[column_index - 1] = max(
                            column_widths[column_index - 1], len(str(parsed_value or ""))
                        )

        if row_count == 0:
            sheet["A1"] = "No results"
            column_widths = [10]

        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions

        for column_index, content_width in enumerate(column_widths, start=1):
            width = min(max(content_width + 2, 11), 42)
            sheet.column_dimensions[get_column_letter(column_index)].width = width


    @staticmethod
    def _populate_empty_summary_sheet(
        sheet, short_telomere_threshold_bp=2000
    ) -> None:
        """Create a valid run summary when no reads passed filtration."""
        sheet.append([
            "barcode",
            "amount_of_telomeres",
            "median_telomere_length",
            f"below_{short_telomere_threshold_bp}bp_pct",
            "med_read_len",
            "mean_density",
            "mean_telo_start",
        ])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions


    @staticmethod
    def _ensure_summary_barcodes(sheet, barcodes: List[str]) -> None:
        """Represent barcodes with zero filtered reads in the central summary."""
        existing = {
            str(sheet.cell(row=row, column=1).value).lower()
            for row in range(2, sheet.max_row + 1)
            if sheet.cell(row=row, column=1).value
        }
        for barcode in barcodes:
            if barcode.lower() not in existing:
                sheet.append([barcode, 0])
        sheet.auto_filter.ref = sheet.dimensions


    @staticmethod
    def _populate_empty_filtered_sheet(sheet, raw_csv: Path) -> None:
        """Create a header-only filtered sheet for a barcode with no passing reads."""
        with raw_csv.open("r", encoding="utf-8-sig", newline="") as handle:
            header = next(csv.reader(handle), [])
        header = ["read_id" if name == "sequence_ID" else name for name in header]
        for derived_name in ("running_median", "seqLen_runningMED", "barcode"):
            if derived_name not in header:
                header.append(derived_name)
        sheet.append(header or ["No reads passed filtering"])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
